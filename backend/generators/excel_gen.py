from pathlib import Path
from typing import List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_styled_excel(
    filepath: str,
    sheet_title: str,
    headers: List[str],
    rows: List[List[Any]],
    theme_color: str = "4F46E5",  # Indigo
    highlight_totals: bool = True,
) -> str:
    """
    Создает красиво стилизованную таблицу Excel с авто-шириной колонок,
    контрастными заголовками, чередующейся подсветкой строк и границами.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Нормализуем цвет (удаляем # для openpyxl)
    clean_theme_color = (theme_color or "4F46E5").lstrip("#").upper()

    # Шрифты и заливки
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=clean_theme_color, end_color=clean_theme_color, fill_type="solid")
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=10, color="1E293B")
    total_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    total_fill = PatternFill(start_color="EEF2F6", end_color="EEF2F6", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    double_bottom_border = Border(
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="double", color="475569"),
    )

    # Запись заголовков
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Запись строк данных
    num_rows = len(rows)
    for row_idx, row_data in enumerate(rows, start=2):
        is_last_row = (row_idx == num_rows + 1) and highlight_totals and (num_rows > 1)
        is_zebra = (row_idx % 2 == 0) and not is_last_row

        ws.row_dimensions[row_idx].height = 22

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border

            # Автоматическое определение выравнивания и числового формата (исключая bool)
            if isinstance(value, bool):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(value, float):
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if is_last_row:
                cell.font = total_font
                cell.fill = total_fill
                cell.border = double_bottom_border
            elif is_zebra:
                cell.font = regular_font
                cell.fill = zebra_fill
            else:
                cell.font = regular_font

    # Авто-подбор ширины колонок
    for col in ws.columns:
        if not col:
            continue
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            lines = val_str.split("\n")
            line_max = max((len(l) for l in lines), default=0)
            max_len = max(max_len, line_max)
        # Ограничиваем разумными пределами (от 12 до 60)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

    # Включаем сетку
    try:
        ws.sheet_view.showGridLines = True
    except Exception:
        if getattr(ws, "views", None) and getattr(ws.views, "sheetView", None) and len(ws.views.sheetView) > 0:
            ws.views.sheetView[0].showGridLines = True

    out_path = Path(filepath)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return str(out_path)
