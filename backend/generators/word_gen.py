from pathlib import Path
from typing import List, Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def set_cell_background(cell, fill_hex: str):
    """Устанавливает цвет фона ячейки таблицы."""
    clean_hex = (fill_hex or "3B82F6").lstrip("#")
    tc_pr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{clean_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Задает внутренние отступы (padding) ячейки."""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def create_styled_docx(
    filepath: str,
    title: str,
    subtitle: Optional[str] = None,
    sections: Optional[List[Dict[str, Any]]] = None,
    theme_color: str = "3B82F6",  # Blue
) -> str:
    """
    Создает профессионально оформленный документ Word с иерархией заголовков,
    аккуратными таблицами и врезками.
    """
    clean_theme_color = (theme_color or "3B82F6").lstrip("#")
    doc = docx.Document()

    # Поля страницы
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    # Титульный заголовок
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(12)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Горизонтальная разделительная черта
    p_hr = doc.add_paragraph()
    p_hr_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{clean_theme_color}"/></w:pBdr>')
    p_hr._element.get_or_add_pPr().append(p_hr_border)

    if sections:
        for sec in sections:
            sec_type = sec.get("type", "paragraph")

            if sec_type == "heading":
                level = sec.get("level", 1)
                h = doc.add_heading(sec.get("text", ""), level=level)
                h.paragraph_format.space_before = Pt(14)
                h.paragraph_format.space_after = Pt(4)
                for r in h.runs:
                    r.font.name = "Calibri"
                    if level == 1:
                        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                    else:
                        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

            elif sec_type == "paragraph":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(sec.get("text", ""))
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

            elif sec_type == "quote":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                # Левая граница цитаты
                border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="{clean_theme_color}"/></w:pBdr>')
                p._element.get_or_add_pPr().append(border)
                r = p.add_run(sec.get("text", ""))
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

            elif sec_type == "bullet":
                items = sec.get("items", [])
                for it in items:
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.space_after = Pt(3)
                    r = bp.add_run(it)
                    r.font.name = "Calibri"
                    r.font.size = Pt(10.5)

            elif sec_type == "table":
                headers = sec.get("headers", [])
                rows = sec.get("rows", [])
                cols_count = max(len(headers), max((len(r) for r in rows), default=0))
                
                if cols_count > 0 and (headers or rows):
                    has_headers = bool(headers)
                    total_table_rows = len(rows) + (1 if has_headers else 0)
                    tbl = doc.add_table(rows=total_table_rows, cols=cols_count)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = True

                    # Заголовки таблицы
                    if has_headers:
                        hdr_cells = tbl.rows[0].cells
                        for idx in range(cols_count):
                            h_text = headers[idx] if idx < len(headers) else ""
                            cell = hdr_cells[idx]
                            set_cell_background(cell, clean_theme_color)
                            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            r = p.add_run(str(h_text))
                            r.font.bold = True
                            r.font.name = "Calibri"
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            r.font.size = Pt(10)

                    # Данные таблицы
                    row_offset = 1 if has_headers else 0
                    for r_idx, r_data in enumerate(rows, start=row_offset):
                        is_zebra = (r_idx % 2 == 0)
                        row_cells = tbl.rows[r_idx].cells
                        for c_idx in range(cols_count):
                            val = r_data[c_idx] if c_idx < len(r_data) else ""
                            cell = row_cells[c_idx]
                            set_cell_margins(cell, top=90, bottom=90, left=160, right=160)
                            if is_zebra:
                                set_cell_background(cell, "F8FAFC")
                            p = cell.paragraphs[0]
                            r = p.add_run(str(val))
                            r.font.name = "Calibri"
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

                    doc.add_paragraph()  # Отступ после таблицы

    out_path = Path(filepath)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return str(out_path)
