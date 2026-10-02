from pathlib import Path
from typing import List, Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


def create_styled_presentation(
    filepath: str,
    title: str,
    subtitle: str = "Подготовлено OmniAgent Local",
    slides_data: Optional[List[Dict[str, Any]]] = None,
    dark_mode: bool = True,
) -> str:
    """
    Создает современную презентацию PowerPoint 16:9.
    Поддерживает:
    - Титульный слайд
    - Слайды со списками тезисов
    - Слайды с карточками метрик / сравнением
    """
    prs = Presentation()
    # Установка пропорций 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # Полностью чистый слайд для ручной компоновки

    # Цветовая палитра
    if dark_mode:
        bg_rgb = RGBColor(0x0F, 0x11, 0x17)
        card_bg = RGBColor(0x1E, 0x22, 0x2D)
        text_primary = RGBColor(0xF8, 0xFA, 0xFC)
        text_secondary = RGBColor(0x94, 0xA3, 0xB8)
        accent_color = RGBColor(0x63, 0x66, 0xF1)  # Indigo
    else:
        bg_rgb = RGBColor(0xF8, 0xFA, 0xFC)
        card_bg = RGBColor(0xFF, 0xFF, 0xFF)
        text_primary = RGBColor(0x0F, 0x17, 0x2A)
        text_secondary = RGBColor(0x64, 0x74, 0x8B)
        accent_color = RGBColor(0x4F, 0x46, 0xE5)

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = bg_rgb
        bg_shape.line.color.rgb = bg_rgb

    # 1. ТИТУЛЬНЫЙ СЛАЙД
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Декоративная полоса акцента
    accent_bar = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.15), Inches(2.8)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent_color
    accent_bar.line.color.rgb = accent_color

    tx_box = slide1.shapes.add_textbox(Inches(1.6), Inches(2.1), Inches(10.5), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.name = "Calibri"
    p1.font.color.rgb = text_primary

    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(20)
    p2.font.name = "Calibri"
    p2.font.color.rgb = text_secondary
    p2.space_before = Pt(16)

    # 2. СЛАЙДЫ С КОНТЕНТОМ
    if slides_data:
        for s_info in slides_data:
            slide = prs.slides.add_slide(blank_layout)
            set_slide_background(slide)

            # Заголовок слайда
            h_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.2))
            htf = h_box.text_frame
            hp = htf.paragraphs[0]
            hp.text = s_info.get("title", "")
            hp.font.size = Pt(28)
            hp.font.bold = True
            hp.font.name = "Calibri"
            hp.font.color.rgb = text_primary

            slide_type = s_info.get("type", "bullets")

            if slide_type == "cards":
                # Несколько карточек (например, 3 колонки)
                cards = s_info.get("cards", [])
                count = min(len(cards), 4)
                if count > 0:
                    margin_left = Inches(1.0)
                    total_width = Inches(11.33)
                    spacing = Inches(0.3)
                    card_width = (total_width - (spacing * (count - 1))) / count

                    for i, card in enumerate(cards[:count]):
                        left = margin_left + i * (card_width + spacing)
                        top = Inches(2.2)
                        card_height = Inches(4.3)

                        card_shape = slide.shapes.add_shape(
                            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_width, card_height
                        )
                        card_shape.fill.solid()
                        card_shape.fill.fore_color.rgb = card_bg
                        card_shape.line.color.rgb = accent_color

                        ctf = card_shape.text_frame
                        ctf.word_wrap = True
                        ctf.margin_top = Inches(0.3)
                        ctf.margin_left = Inches(0.3)
                        ctf.margin_right = Inches(0.3)

                        # Название карточки
                        cp1 = ctf.paragraphs[0]
                        cp1.text = card.get("title", "")
                        cp1.font.bold = True
                        cp1.font.size = Pt(20)
                        cp1.font.color.rgb = text_primary

                        # Текст карточки
                        cp2 = ctf.add_paragraph()
                        cp2.text = card.get("text", "")
                        cp2.font.size = Pt(14)
                        cp2.font.color.rgb = text_secondary
                        cp2.space_before = Pt(12)

            else:
                # Стандартный слайд со списком пунктов
                bullets = s_info.get("bullets", [])
                b_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
                btf = b_box.text_frame
                btf.word_wrap = True

                for b_idx, bullet_text in enumerate(bullets):
                    bp = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
                    bp.text = f"•   {bullet_text}"
                    bp.font.size = Pt(18)
                    bp.font.name = "Calibri"
                    bp.font.color.rgb = text_primary
                    bp.space_before = Pt(16)
                    bp.line_spacing = 1.2

    out_path = Path(filepath)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return str(out_path)
